import os
import time
import json
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client=Groq(api_key=my_api_key)
model = "openai/gpt-oss-20b"


class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    other_requirements: list[str]
    responsibilities: list[str]

jobd_schema = JobD.model_json_schema()

def parse_job_description(job_description):
    system_prompt = f"""
    You are an expert HR assistant.

    Your job is to analyze job descriptions and extract
    structured information from them.

    Return ONLY valid JSON matching this schema:

    {jobd_schema}

    IMPORTANT:
    Do NOT return the schema itself.
    Do NOT return fields like "properties", "title" or "type".
    Fill the schema with actual information extracted from the job description.

    If minimum experience is not mentioned, return null.
    If information for a list is missing, return an empty list.
    Do not invent information.

    IMPORTANT CLASSIFICATION RULES:

    1. required_skills:
        Only include technical skills, technologies, tools,
        frameworks, libraries, databases, programming languages,
        platforms, or clearly defined technical abilities that are
        explicitly required for the role.

        Do NOT include generic soft skills such as:
        communication, teamwork, leadership, problem-solving,
        analytical thinking, adaptability, or time management.

    2. preferred_skills:
        Only include technical skills, technologies, tools,
        frameworks, libraries, databases, programming languages,
        platforms, or clearly defined technical abilities that are
        preferred or considered a plus.

        Do NOT include generic soft skills such as:
        communication, teamwork, leadership, problem-solving,
        analytical thinking, adaptability, or time management.

        Do NOT include vague phrases such as:
        "debugging skills", "troubleshooting ability", or
        "understanding of software development" unless they refer
        to a clearly defined technical technology, tool, process,
        or domain.

    3. education_requirements:
       Include degree, academic qualification, field of study,
       or educational background requirements.

    4. other_requirements:
       Include requirements that are not skills or education,
       such as age requirements, work authorization, location
       requirements, certifications, or other eligibility conditions.

    5. responsibilities:
       Include the actual duties and responsibilities of the role.

    Do NOT put education requirements or general eligibility
    requirements inside required_skills or preferred_skills.
    """

    user_prompt = f"""
    Analyze the following job description:

    {job_description}
    """

    message_system = {
        "role": "system",
        "content": system_prompt
    }

    message_user = {
        "role": "user",
        "content": user_prompt
    }

    messages = [message_system, message_user]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    raw_json = response.choices[0].message.content

    job_data = json.loads(raw_json)

    return JobD(**job_data)



#parse real
class MatchDetails(BaseModel):
    candidate_name: str | None = None
    matching_skills: list[str] = Field(default_factory=list)
    missing_important_skills: list[str] = Field(default_factory=list)
    experience_requirement_met: bool
    education_requirement_met: bool
    other_requirements_met: bool
    final_verdict: str

class ScoreBreakdown(BaseModel):
    required_skills: float
    experience: float
    preferred_skills: float
    education: float
    other_requirements: float

class MatchResult(BaseModel):
    score: float | None = None
    breakdown: ScoreBreakdown | None = None
    details: MatchDetails

class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


resume_schema = Resume.model_json_schema()

def normalize_skill(skill):
    skill = skill.lower().strip()

    aliases = {
        "js": "javascript",
        "ts": "typescript",
        "react.js": "react",
        "reactjs": "react",
        "node.js": "node",
        "nodejs": "node",
        "postgres": "postgresql",
        "postgre": "postgresql",
        "mongo": "mongodb",
        "oop": "object oriented programming",
        "object-oriented programming": "object oriented programming",
        "dsa": "data structures and algorithms",
    }

    return aliases.get(skill, skill)


def skill_match(required_skill, resume_skills):
    required_skill = normalize_skill(required_skill)

    for resume_skill in resume_skills:
        if required_skill == normalize_skill(resume_skill):
            return True

    return False

def calculate_score(job, resume, match_details):
    required_score = 0.0
    experience_score = 0.0
    preferred_score = 0.0
    education_score = 0.0
    other_requirements_score = 0.0

    # Required skills — 50%
    required_skills = job.required_skills
    resume_skills = resume.skills

    if required_skills:
        required_matches = sum(
            skill_match(skill, resume_skills)
            for skill in required_skills
        )

        required_match_ratio = required_matches / len(required_skills)
        required_score = required_match_ratio * 50

    # Experience — 25%
    if job.minimum_experience is None:
        experience_score = 25
    elif match_details.experience_requirement_met:
        experience_score = 25

    # Preferred skills — 10%
    preferred_skills = job.preferred_skills

    if preferred_skills:
        preferred_matches = sum(
            skill_match(skill, resume_skills)
            for skill in preferred_skills
        )

        preferred_match_ratio = preferred_matches / len(preferred_skills)
        preferred_score = preferred_match_ratio * 10

    # Education — 10%
    if match_details.education_requirement_met:
        education_score = 10

    # Other requirements — 5%
    if match_details.other_requirements_met:
        other_requirements_score = 5

    score = (
    required_score
    + experience_score
    + preferred_score
    + education_score
    + other_requirements_score
)

    breakdown = ScoreBreakdown(
        required_skills=round(required_score, 2),
        experience=round(experience_score, 2),
        preferred_skills=round(preferred_score, 2),
        education=round(education_score, 2),
        other_requirements=round(other_requirements_score, 2),
    )

    return round(score, 2), breakdown

def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()

    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}

    Return ONLY valid JSON matching this schema:

    {match_schema}

    IMPORTANT:
    - Return actual analysis data, NOT the schema itself.
    - "details" MUST contain exactly these fields:
        candidate_name
        matching_skills
        missing_important_skills
        experience_requirement_met
        education_requirement_met
        other_requirements_met
        final_verdict

    Rules:

    1. candidate_name:
       Use the candidate's name from the resume.

    2. matching_skills:
       List skills from the candidate that match the job requirements.

    3. missing_important_skills:
       List important skills required by the job that are missing
       from the candidate's resume.

    4. experience_requirement_met:
       Return true if the candidate satisfies the job's minimum
       experience requirement.
       Otherwise return false.

    5. education_requirement_met:
        Return true if the candidate satisfies the job's education
        requirements. Otherwise return false.
        If the job has no education requirements, return true.

    6. final_verdict:
        Give a short 1-2 sentence hiring assessment.

    7. other_requirements_met:
        Return true if the candidate satisfies the job's other
        requirements. Otherwise return false.
        If the job has no other requirements, return true.

    8. Do not invent information.

    9. Do NOT calculate or return a score. Only provide the fields inside "details".

    10. Return ONLY JSON.
    """

    message = {
        "role": "user",
        "content": prompt
    }

    messages = [message]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    raw_output = response.choices[0].message.content

    data = json.loads(raw_output)

    match_result = MatchResult(**data)

    score, breakdown = calculate_score(job, resume, match_result.details)

    return MatchResult(
        score=score,
        breakdown=breakdown,
        details=match_result.details
    )

def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



